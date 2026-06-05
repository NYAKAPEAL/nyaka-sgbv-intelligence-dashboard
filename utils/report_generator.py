"""
utils/report_generator.py
Generates a downloadable ~5-page monthly programmatic report (Word .docx)
from the live dashboard data, in Nyaka brand styling.

Design intent: produces a solid DRAFT that the programme lead edits and
adds their own voice to before sending to funders — not a final auto-document.
"""
import io
from datetime import datetime

import pandas as pd
import numpy as np

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

# ── Nyaka brand palette ───────────────────────────────────────────────────────
PURPLE   = RGBColor(0x5B, 0x2D, 0x8E)
PURPLE_DK= RGBColor(0x3C, 0x1A, 0x6B)
ORANGE   = RGBColor(0xE0, 0x7B, 0x29)
GOLD     = RGBColor(0xC9, 0xA8, 0x4C)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
RED      = RGBColor(0xC6, 0x28, 0x28)
GREY     = RGBColor(0x54, 0x6E, 0x7A)
DARK     = RGBColor(0x1A, 0x1A, 0x2E)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

PALETTE_HEX = ["#5B2D8E", "#E07B29", "#C9A84C", "#2E7D32", "#00695C", "#C62828"]


# ── chart helpers (matplotlib → PNG bytes) ────────────────────────────────────
def _bar_png(labels, values, title, color="#5B2D8E", horizontal=False):
    fig, ax = plt.subplots(figsize=(6.2, 2.6), dpi=150)
    if horizontal:
        ax.barh(labels, values, color=color)
        ax.invert_yaxis()
    else:
        ax.bar(labels, values, color=color)
        plt.xticks(rotation=30, ha="right", fontsize=7)
    ax.set_title(title, fontsize=9, fontweight="bold", color="#3C1A6B")
    ax.tick_params(labelsize=7)
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _donut_png(labels, values, title):
    fig, ax = plt.subplots(figsize=(3.4, 2.6), dpi=150)
    ax.pie(values, labels=labels, autopct="%1.0f%%", startangle=90,
           colors=PALETTE_HEX[:len(values)],
           wedgeprops=dict(width=0.42, edgecolor="white"),
           textprops=dict(fontsize=7))
    ax.set_title(title, fontsize=9, fontweight="bold", color="#3C1A6B")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


# ── docx low-level helpers ────────────────────────────────────────────────────
def _shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=9, align=None, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    if white:
        run.font.color.rgb = WHITE
    elif color is not None:
        run.font.color.rgb = color


def _heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = PURPLE
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "E07B29")
        pbdr.append(bottom)
        pPr.append(pbdr)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = ORANGE
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def _body(doc, text, size=10, color=DARK, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(4)
    return p


def _kpi_table(doc, kpis):
    """kpis: list of (label, value, sublabel) tuples. Renders a branded grid."""
    n = len(kpis)
    cols = min(n, 3)
    rows = (n + cols - 1) // cols
    table = doc.add_table(rows=rows * 2, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (label, value, sub) in enumerate(kpis):
        r = (idx // cols) * 2
        c = idx % cols
        val_cell = table.cell(r, c)
        _shade_cell(val_cell, "5B2D8E")
        _set_cell_text(val_cell, value, bold=True, size=18, white=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
        lbl_cell = table.cell(r + 1, c)
        _shade_cell(lbl_cell, "EDE7F6")
        _set_cell_text(lbl_cell, f"{label}\n{sub}" if sub else label,
                       bold=False, size=8, color=PURPLE_DK,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
    # widths
    w = int(9360 / cols)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(w / 1440)
    return table


def _data_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _shade_cell(hdr[i], "5B2D8E")
        _set_cell_text(hdr[i], h, bold=True, size=9, white=True)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "F5F2FC" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row):
            _shade_cell(cells[ci], fill)
            _set_cell_text(cells[ci], val, size=9, color=DARK)
    if col_widths:
        for row in table.rows:
            for ci, wdt in enumerate(col_widths):
                row.cells[ci].width = Inches(wdt)
    return table


# ── data filtering ────────────────────────────────────────────────────────────
def _month_window(df, date_col, year, month):
    d = pd.to_datetime(df.get(date_col), errors="coerce")
    return df[(d.dt.year == year) & (d.dt.month == month)]


# ── main report builder ───────────────────────────────────────────────────────
def build_monthly_report(data: dict, year: int, month: int,
                          prepared_by: str = "") -> io.BytesIO:
    """
    data keys expected (any missing → section gracefully notes 'no data'):
      survivors, perpetrators, perp_followup, outreach, school, narrative
    Returns BytesIO of the .docx.
    """
    month_name = datetime(year, month, 1).strftime("%B %Y")
    doc = Document()

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # Page size US Letter, 1" margins
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1))

    # ── COVER BLOCK ──
    # Nyaka logo, centred
    import os as _os
    _logo = _os.path.join(_os.path.dirname(__file__), "..", "assets", "logo",
                          "nyaka_horizontal.png")
    if _os.path.exists(_logo):
        lp = doc.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            lp.add_run().add_picture(_logo, width=Inches(2.6))
        except Exception:
            pass

    org = doc.add_paragraph()
    org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = org.add_run("NYAKA — Sexual & Gender-Based Violence Programme")
    run.font.name = "Arial"; run.bold = True; run.font.size = Pt(13)
    run.font.color.rgb = PURPLE

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Kanungu · Rukungiri · Rubanda Districts, South-Western Uganda")
    r2.font.name = "Arial"; r2.font.size = Pt(9); r2.font.color.rgb = GREY

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = title.add_run(f"Monthly Programmatic Report")
    rt.font.name = "Arial"; rt.bold = True; rt.font.size = Pt(20); rt.font.color.rgb = DARK
    title.paragraph_format.space_before = Pt(10)

    per = doc.add_paragraph()
    per.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp = per.add_run(month_name)
    rp.font.name = "Arial"; rp.bold = True; rp.font.size = Pt(14); rp.font.color.rgb = ORANGE

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run(f"Generated {datetime.now():%d %B %Y}"
                      + (f"  ·  Prepared by {prepared_by}" if prepared_by else ""))
    rm.font.name = "Arial"; rm.font.size = Pt(8); rm.font.color.rgb = GREY
    meta.paragraph_format.space_after = Pt(8)

    # Draft watermark notice
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run("DRAFT — review, verify figures, and add programme narrative "
                      "before sharing externally.")
    rn.font.name = "Arial"; rn.italic = True; rn.font.size = Pt(8); rn.font.color.rgb = RED

    # ── gather monthly slices ──
    sv = data.get("survivors", pd.DataFrame())
    pp = data.get("perpetrators", pd.DataFrame())
    ppf = data.get("perp_followup", pd.DataFrame())
    ot = data.get("outreach", pd.DataFrame())
    sc = data.get("school", pd.DataFrame())

    sv_m = _month_window(sv[sv.get("report_category", pd.Series("", index=sv.index)).str.contains("enrollment", na=False)]
                         if "report_category" in sv.columns else sv,
                         "date_enrolled", year, month) if len(sv) else pd.DataFrame()
    pp_e = pp[pp.get("report_category", pd.Series("", index=pp.index)) == "enrollment"] if "report_category" in pp.columns else pp
    pp_m = _month_window(pp_e, "date_enrolled", year, month) if len(pp_e) else pd.DataFrame()
    ppf_m = _month_window(ppf, "follow_up_date", year, month) if len(ppf) else pd.DataFrame()
    ot_m = _month_window(ot, "date", year, month) if len(ot) else pd.DataFrame()
    sc_m = _month_window(sc, "date", year, month) if len(sc) else pd.DataFrame()

    # ── 1. EXECUTIVE SUMMARY ──
    _heading(doc, "1. Executive Summary", 1)

    n_surv = len(sv_m)
    n_perp = len(pp_m)
    arrests = 0
    if len(pp_m) and "perpetrator_location_label" in pp_m.columns:
        arrests = int(pp_m["perpetrator_location_label"].astype(str)
                      .str.contains("Arrest", na=False).sum())
    n_out = len(ot_m)
    reach = int(ot_m["total"].sum()) if len(ot_m) and "total" in ot_m.columns else 0
    n_sch = len(sc_m)
    students = int(sc_m["total"].sum()) if len(sc_m) and "total" in sc_m.columns else 0

    # Narrative summary (auto-drafted, factual)
    female_pct = ""
    if len(sv_m) and "client_gender" in sv_m.columns:
        fp = (sv_m["client_gender"].astype(str).str.upper().str[0] == "F").mean() * 100
        female_pct = f" Of these, {fp:.0f}% were female."
    top_dist = ""
    if len(sv_m) and "district" in sv_m.columns:
        td = sv_m["district"].value_counts()
        if len(td):
            top_dist = f" The highest case volume was in {td.index[0]} ({td.iloc[0]} survivors)."

    _body(doc,
          f"During {month_name}, the Nyaka SGBV programme supported {n_surv} newly "
          f"enrolled survivors across its three operational districts.{female_pct}{top_dist} "
          f"The legal team registered {n_perp} perpetrator cases, of which {arrests} "
          f"resulted in arrest during the month. Prevention work reached an estimated "
          f"{reach:,} community members through {n_out} outreach sessions, and {n_sch} "
          f"school engagement sessions reached {students:,} students.")

    _body(doc,
          "This report draws directly from field data captured through the programme's "
          "SurveyCTO tools and the SGBV monitoring dashboard. Figures should be verified "
          "against source records before external use.",
          size=9, italic=True, color=GREY)

    # KPI grid
    _kpi_table(doc, [
        ("Survivors supported", f"{n_surv}", month_name),
        ("Perpetrator cases", f"{n_perp}", f"{arrests} arrested"),
        ("Outreach sessions", f"{n_out}", f"{reach:,} reached"),
        ("School sessions", f"{n_sch}", f"{students:,} students"),
        ("Arrests made", f"{arrests}", "this month"),
        ("Districts active", "3", "Kanungu·Rukungiri·Rubanda"),
    ])

    # ── 2. SURVIVOR SUPPORT ──
    _heading(doc, "2. Survivor Support", 1)
    if len(sv_m):
        _body(doc, f"{n_surv} survivors were enrolled in {month_name}. The breakdown by "
                   "district and violence type is shown below.")
        # District table
        if "district" in sv_m.columns:
            dist_counts = sv_m["district"].value_counts()
            rows = [[str(d), str(c)] for d, c in dist_counts.items()]
            _data_table(doc, ["District", "Survivors enrolled"], rows, [3.5, 2.5])
            # chart
            img = _bar_png(list(dist_counts.index), list(dist_counts.values),
                           f"Survivors by district — {month_name}", color="#5B2D8E")
            doc.add_picture(img, width=Inches(5.5))
        # Violence type
        vcol = next((c for c in ["assault_label", "violence_type", "sgbv_type_label"]
                     if c in sv_m.columns), None)
        if vcol:
            vt = sv_m[vcol].value_counts().head(5)
            if len(vt):
                _heading(doc, "Violence types reported", 2)
                rows = [[str(v), str(c)] for v, c in vt.items()]
                _data_table(doc, ["Violence type", "Cases"], rows, [4.0, 2.0])
    else:
        _body(doc, f"No survivor enrollments were recorded in {month_name}.", italic=True, color=GREY)

    # ── 3. JUSTICE & LEGAL PIPELINE ──
    _heading(doc, "3. Justice & Legal Pipeline", 1)
    if len(pp_m) or len(ppf_m):
        _body(doc, f"{n_perp} new perpetrator cases were registered in {month_name}, "
                   f"with {arrests} arrests recorded. Case status movements logged during "
                   f"the month are summarised below.")
        if len(ppf_m) and "case_status_overview" in ppf_m.columns:
            status = ppf_m["case_status_overview"].value_counts().head(6)
            rows = [[str(s), str(c)] for s, c in status.items()]
            _data_table(doc, ["Case status", "Follow-up records"], rows, [4.0, 2.0])
            img = _bar_png(list(status.index), list(status.values),
                           f"Case status activity — {month_name}",
                           color="#E07B29", horizontal=True)
            doc.add_picture(img, width=Inches(5.5))
    else:
        _body(doc, f"No new perpetrator cases or legal follow-up were recorded in {month_name}.",
              italic=True, color=GREY)

    # ── 4. PREVENTION & OUTREACH ──
    _heading(doc, "4. Prevention & Outreach", 1)
    if len(ot_m) or len(sc_m):
        _body(doc, f"The prevention team delivered {n_out} community outreach sessions "
                   f"(reaching approximately {reach:,} people) and {n_sch} school "
                   f"engagement sessions (reaching {students:,} students) during {month_name}.")
        prev_rows = [
            ["Community outreach sessions", f"{n_out}"],
            ["People reached (outreach)", f"{reach:,}"],
            ["School sessions", f"{n_sch}"],
            ["Students reached", f"{students:,}"],
        ]
        _data_table(doc, ["Prevention indicator", "This month"], prev_rows, [4.5, 2.0])
    else:
        _body(doc, f"No prevention or outreach activity was recorded in {month_name}.",
              italic=True, color=GREY)

    # ── 5. SAFEGUARDING & CASE FOLLOW-UP ──
    _heading(doc, "5. Safeguarding & Case Follow-up", 1)
    narrative = data.get("narrative", pd.DataFrame())
    stale_n, nofu_n = None, None
    if len(narrative) and "ageing_band" in narrative.columns:
        open_n = narrative[~narrative.get("is_closed", False)] if "is_closed" in narrative.columns else narrative
        stale_n = int((open_n["ageing_band"] == "180+ days (stale)").sum())
        if "last_followup_date" in narrative.columns:
            nofu_n = int(narrative["last_followup_date"].isna().sum())
    if stale_n is not None:
        _body(doc,
              f"As of this report, {stale_n} open cases across the whole programme "
              f"(all enrolment periods, not only {month_name}) have had no follow-up "
              f"update in over 180 days and require attention. "
              + (f"{nofu_n} enrolled cases have no recorded follow-up visit yet. "
                 if nofu_n else "")
              + "The case ageing report in the dashboard provides the full priority worklist.")
        _body(doc,
              "Safeguarding note: all survivor data in this report is anonymised. "
              "Crisis-flagged cases are managed separately through the clinical review "
              "process and are not individually identified here.",
              size=9, italic=True, color=GREY)
    else:
        _body(doc, "Case ageing data was not available for this reporting period. "
                   "See the dashboard Narratives page for live follow-up tracking.",
              italic=True, color=GREY)

    # ── 6. CHALLENGES, RECOMMENDATIONS & NEXT STEPS ──
    _heading(doc, "6. Challenges, Recommendations & Next Steps", 1)
    _body(doc, "Challenges (to be completed by programme team):", bold=True, size=10)
    for placeholder in ["[ Add key operational or contextual challenges this month ]",
                        "[ Note any cases requiring escalation or external support ]"]:
        p = doc.add_paragraph(placeholder, style=None)
        p.paragraph_format.left_indent = Inches(0.3)
        for r in p.runs:
            r.font.name = "Arial"; r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True

    _body(doc, "Recommendations & next steps:", bold=True, size=10)
    for placeholder in ["[ Add programme recommendations for next month ]",
                        "[ Note follow-up actions, resourcing needs, or referrals ]"]:
        p = doc.add_paragraph(placeholder)
        p.paragraph_format.left_indent = Inches(0.3)
        for r in p.runs:
            r.font.name = "Arial"; r.font.size = Pt(9); r.font.color.rgb = GREY; r.italic = True

    # Footer line
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = foot.add_run("Nyaka SGBV Programme · Confidential · Generated from the SGBV "
                      "Monitoring Dashboard · Figures to be verified before external use")
    rf.font.name = "Arial"; rf.font.size = Pt(7.5); rf.font.color.rgb = GREY
    foot.paragraph_format.space_before = Pt(16)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return _fix_settings(buf)


def _fix_settings(buf: io.BytesIO) -> io.BytesIO:
    """python-docx emits a <w:zoom> with no percent attr, which fails strict
    OOXML validation. Patch it in-place so the file is fully spec-compliant."""
    import zipfile
    buf.seek(0)
    src = zipfile.ZipFile(buf, "r")
    out_buf = io.BytesIO()
    out = zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED)
    for item in src.namelist():
        content = src.read(item)
        if item == "word/settings.xml":
            text = content.decode("utf-8")
            import re
            if "<w:zoom" in text and "w:percent" not in text:
                # python-docx emits <w:zoom w:val="bestFit"/> with no percent attr.
                text = re.sub(r"<w:zoom([^>]*?)/>",
                              r'<w:zoom\1 w:percent="100"/>', text)
            content = text.encode("utf-8")
        out.writestr(item, content)
    out.close()
    src.close()
    out_buf.seek(0)
    return out_buf